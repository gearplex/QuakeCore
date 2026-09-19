#!/usr/bin/env python3
"""Read-only audit of privately supplied conventional-wall design workbooks.

Requires openpyxl. Writes private JSON/Markdown; never modifies source files or
certifies structural designs. Run against an extracted source package.
"""
import argparse
import hashlib
import json
from pathlib import Path

from openpyxl import load_workbook


FIELDS = {
    5: "stories", 6: "story_height_ft", 7: "fc_ksi", 8: "fy_ksi",
    9: "design_drift_limit", 10: "R", 11: "Cd", 12: "Ss_g", 13: "S1_g",
    14: "wall_length_ft", 15: "wall_thickness_in", 16: "amplified_design_drift",
    17: "Pu_kip", 18: "Vu_kip", 19: "Mu_kip_in", 20: "Ve_kip",
    21: "boundary_As_in2", 22: "boundary_length_in", 35: "phi_Mn_kip_in",
    36: "flexural_DCR", 37: "archetype_id", 38: "design_story",
}


def read_book(path):
    book = load_workbook(path, read_only=True, data_only=True)
    sheets = {}
    for sheet in book:
        if not sheet.title.startswith("Sheet"):
            continue
        rows = list(sheet.iter_rows(min_row=5, max_row=38, min_col=4, max_col=87, values_only=True))
        sheets[sheet.title] = rows
    first = sheets["Sheet1"]
    records = []
    for col in range(84):
        row = {field: first[r-5][col] for r, field in FIELDS.items()}
        if row["archetype_id"] is None:
            continue
        row["source_column_number"] = col + 4
        records.append(row)
    ids = [r["archetype_id"] for r in records]
    if len(ids) != len(set(ids)):
        raise ValueError(f"duplicate archetype ID within {path.name}")
    book.close()
    return {"filename": path.name, "sha256": hashlib.sha256(path.read_bytes()).hexdigest(),
            "archetype_count": len(records), "records": records}, sheets


def main():
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("--source-dir", type=Path, required=True)
    ap.add_argument("--out-dir", type=Path, required=True)
    args = ap.parse_args()
    books, grids = {}, {}
    for path in sorted(args.source_dir.glob("DesignSummary*.xlsx")):
        books[path.name], grids[path.name] = read_book(path)
    bn = "DesignSummary_Special_Bearing_R=6.xlsx"
    nn = "DesignSummary_Special_Nonbearing_R=6.xlsx"
    b = {r["archetype_id"]: r for r in books[bn]["records"]}
    n = {r["archetype_id"]: r for r in books[nn]["records"]}
    if set(b) != set(n):
        raise ValueError("bearing/nonbearing archetype ID sets differ")
    equal_axial = [i for i in b if b[i]["Pu_kip"] is not None and b[i]["Pu_kip"] == n[i]["Pu_kip"]]
    differences = []
    for name, br in grids[bn].items():
        nr = grids[nn][name]
        for ri, (brow, nrow) in enumerate(zip(br, nr), 5):
            for ci, (bv, nv) in enumerate(zip(brow, nrow), 4):
                if bv != nv:
                    differences.append({"sheet": name, "row": ri, "column": ci,
                                        "bearing_value": bv, "nonbearing_value": nv})
    report = {
        "scope": "source audit only; cached workbook values, no recalculation or structural design approval",
        "archetype_key": "workbook variant + archetype_id (IDs are reused across variants)",
        "workbooks": books,
        "R6_first_story_matching_Pu_ids": equal_axial,
        "R6_design_block_differences": differences,
        "generated_all_story_model_input_files": [str(p.relative_to(args.source_dir)) for p in args.source_dir.rglob("AllStoriesOpenseesModelINPUT*")],
        "qualification_ready": False,
        "blocking_findings": [
            "R=6 bearing/nonbearing gravity design must be reconciled before a bearing-wall qualification run.",
            "IDA success must check solver status, final time and valid EDP rows, not only output-file existence.",
            "A 5% drift stop is a screening limit until its relationship to physical collapse is justified.",
            "Published/cached P-695 tables need result-to-model provenance before reuse.",
        ],
    }
    args.out_dir.mkdir(parents=True, exist_ok=True)
    (args.out_dir/"conventional_input_audit.json").write_text(json.dumps(report, indent=2, allow_nan=False)+"\n")
    lines = ["# Conventional wall input audit", "", report["scope"], "",
             "| Workbook | Archetypes | R values | Cd values |", "|---|---:|---|---|"]
    for name, book in books.items():
        values = book["records"]
        lines.append(f"| {name} | {len(values)} | {sorted(set(r['R'] for r in values))} | {sorted(set(r['Cd'] for r in values))} |")
    lines += ["", f"The two R=6 files have identical first-story Pu values for {len(equal_axial)}/{len(b)} IDs.",
              f"Across their D5:CI38 design blocks, only {len(differences)} cells differ.", "",
              "| Sheet | Row | Column number | Bearing value | Nonbearing value |", "|---|---:|---:|---:|---:|"]
    for d in differences:
        lines.append(f"| {d['sheet']} | {d['row']} | {d['column']} | {d['bearing_value']} | {d['nonbearing_value']} |")
    lines += ["", "## Required resolution", ""] + [f"- {s}" for s in report["blocking_findings"]]
    lines += ["", "This audit leaves every original workbook unchanged."]
    (args.out_dir/"conventional_input_audit.md").write_text("\n".join(lines)+"\n")
    print(json.dumps({"workbooks": len(books), "matching_Pu": len(equal_axial),
                      "design_block_differences": len(differences), "qualification_ready": False}))


if __name__ == "__main__":
    main()
